from pathlib import Path
from datetime import datetime, date
import os
import json

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(os.environ.get("QH_PROJECT_ROOT", Path.cwd())).expanduser().resolve()
SOURCE = ROOT / "0718开始分析" / "青海省16-24年气象数据.xlsx"
PAYLOAD = ROOT / "_codex_weather_merge" / "merge_values.json"
OUTPUT = ROOT / "0718开始分析" / "青海省16-24年气象数据_更新平均气温气压.xlsx"
REPORT = ROOT / "_codex_weather_merge" / "final_validation.json"

with PAYLOAD.open("r", encoding="utf-8") as f:
    payload = json.load(f)

new_values = payload["values"]
if len(new_values) != 147960:
    raise ValueError(f"Expected 147960 merged rows, found {len(new_values)}")

source_wb = load_workbook(SOURCE, read_only=True, data_only=False)
source_ws = source_wb["Qinghai_PREC_Long"]

output_wb = Workbook(write_only=True)
output_ws = output_wb.create_sheet("Qinghai_PREC_Long")
output_ws.freeze_panes = "A2"
output_ws.sheet_view.showGridLines = True

widths = {
    "A": 11, "B": 11, "C": 11,
    "D": 18, "E": 23, "F": 11,
    "G": 13, "H": 9,
    "I": 13, "J": 13, "K": 13, "L": 13, "M": 13,
    "N": 13, "O": 13,
}
for col, width in widths.items():
    output_ws.column_dimensions[col].width = width

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
header_alignment = Alignment(horizontal="center", vertical="center")

source_rows = source_ws.iter_rows(values_only=True)
source_header = list(next(source_rows))
expected_source_header = [
    "CountyC", "CityC", "ProvinceC", "County", "City", "Province",
    "date", "year", "降水量", "相对湿度", "表面气压", "风速", "温度",
]
if source_header != expected_source_header:
    raise ValueError(f"Unexpected source header: {source_header}")

headers = source_header + ["平均气温", "平均气压"]
header_cells = []
for value in headers:
    cell = WriteOnlyCell(output_ws, value=value)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    header_cells.append(cell)
output_ws.append(header_cells)

written = 0
for written, (source_row, pair) in enumerate(zip(source_rows, new_values, strict=True), start=1):
    row = list(source_row) + pair

    # Preserve a true Excel date and explicit display format.
    date_cell = WriteOnlyCell(output_ws, value=row[6])
    date_cell.number_format = "yyyy-mm-dd"
    row[6] = date_cell

    # Use a consistent precision for meteorological measurements.
    for idx in range(8, 15):
        cell = WriteOnlyCell(output_ws, value=row[idx])
        cell.number_format = "0.0000"
        row[idx] = cell

    output_ws.append(row)

if written != 147960:
    raise ValueError(f"Expected to write 147960 rows, wrote {written}")

output_ws.auto_filter.ref = f"A1:O{written + 1}"
output_ws.sheet_properties.pageSetUpPr.fitToPage = True
output_ws.page_setup.fitToWidth = 1
output_ws.page_setup.fitToHeight = 0

output_wb.save(OUTPUT)
source_wb.close()

# Independent full-workbook verification.
check_wb = load_workbook(OUTPUT, read_only=True, data_only=True)
check_ws = check_wb["Qinghai_PREC_Long"]
rows = check_ws.iter_rows(values_only=True)
header_check = list(next(rows))

row_count = 0
county_codes = set()
keys = set()
duplicate_keys = 0
missing_temperature = 0
missing_pressure = 0
date_min = None
date_max = None
first_row = None
last_row = None
max_abs_temp_diff = 0.0
max_abs_pres_diff = 0.0

for idx, row in enumerate(rows):
    row_count += 1
    if first_row is None:
        first_row = list(row)
    last_row = list(row)
    county = int(row[0])
    dt = row[6]
    if isinstance(dt, datetime):
        dt = dt.date()
    if not isinstance(dt, date):
        raise TypeError(f"Row {idx + 2}: date is not an Excel date: {dt!r}")
    county_codes.add(county)
    key = (county, dt)
    if key in keys:
        duplicate_keys += 1
    keys.add(key)
    date_min = dt if date_min is None or dt < date_min else date_min
    date_max = dt if date_max is None or dt > date_max else date_max

    temp = row[13]
    pres = row[14]
    if temp is None:
        missing_temperature += 1
    else:
        max_abs_temp_diff = max(max_abs_temp_diff, abs(float(temp) - float(new_values[idx][0])))
    if pres is None:
        missing_pressure += 1
    else:
        max_abs_pres_diff = max(max_abs_pres_diff, abs(float(pres) - float(new_values[idx][1])))

check_wb.close()

validation = {
    "output": str(OUTPUT),
    "sheet_names": ["Qinghai_PREC_Long"],
    "headers": header_check,
    "data_rows": row_count,
    "columns": len(header_check),
    "unique_counties": len(county_codes),
    "unique_county_date_keys": len(keys),
    "duplicate_county_date_keys": duplicate_keys,
    "date_min": date_min.isoformat(),
    "date_max": date_max.isoformat(),
    "missing_平均气温": missing_temperature,
    "missing_平均气压": missing_pressure,
    "max_abs_difference_平均气温": max_abs_temp_diff,
    "max_abs_difference_平均气压": max_abs_pres_diff,
    "first_data_row": first_row,
    "last_data_row": last_row,
    "file_size_bytes": OUTPUT.stat().st_size,
}

with REPORT.open("w", encoding="utf-8") as f:
    json.dump(validation, f, ensure_ascii=False, indent=2, default=str)

print(json.dumps(validation, ensure_ascii=False, indent=2, default=str))
